#!/usr/bin/env python3
"""Apply the approved manuscript-wide chapter-role title audit safely.

The approved audit lives on editor/chapter-title-role-audit-batch-001. This tool
reconciles that frozen audit against the current checkout before changing titles.
It changes title metadata/headings only. Prose is never rewritten.

Requires python-docx when applying/verifying Book I/II canonical DOCX sources.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import subprocess
import tempfile
from pathlib import Path
from typing import Iterable

AUDIT_REF = "origin/editor/chapter-title-role-audit-batch-001"
AUDIT_FILES = [
    "state/CHAPTER_TITLE_ROLE_AUDIT.md",
    "state/chapter-title-role-audit/BATCH_081_120.md",
    "state/chapter-title-role-audit/BATCH_121_160.md",
    "state/chapter-title-role-audit/BATCH_161_200.md",
    "state/chapter-title-role-audit/BATCH_201_240.md",
    "state/chapter-title-role-audit/BATCH_241_280.md",
    "state/chapter-title-role-audit/BATCH_281_320.md",
    "state/chapter-title-role-audit/BATCH_321_360.md",
    "state/chapter-title-role-audit/BATCH_361_400.md",
    "state/chapter-title-role-audit/BATCH_401_440.md",
    "state/chapter-title-role-audit/BATCH_441_488.md",
]
FINAL_APPROVAL = "state/chapter-title-role-audit/FINAL_APPROVAL.md"
EXPECTED_RENAMES = 118
RETIRED_CHAPTER = 150

BOOK1_DOCX = Path("state/manuscript/Peg_Leg_Greg_authoritative_ch82_final_name_map.docx")
BOOK2_DOCX = Path("state/manuscript/Peg_Leg_Greg_Book2_Manuscript_Ch83-137.docx")
RUN_138_155 = Path("state/manuscript/Peg_Leg_Greg_Running_Manuscript_Ch138-155.md")
RECOVERED_156_219 = Path("state/manuscript/Peg_Leg_Greg_Recovered_Ch156-219_EXACT.md")
RUNNING_220 = Path("state/manuscript/Peg_Leg_Greg_Running_Manuscript.md")
CHAPTER_INDEX = Path("state/MANUSCRIPT_CHAPTER_INDEX.md")


def normalize_title(value: str) -> str:
    value = html.unescape(re.sub(r"<[^>]+>", "", value))
    return re.sub(r"\s+", " ", value).strip().casefold()


def clean_markdown_cell(value: str) -> str:
    value = value.strip()
    value = value.replace("`", "").replace("**", "")
    return value.strip()


def parse_audit_table(text: str) -> dict[int, tuple[str, str]]:
    result: dict[int, tuple[str, str]] = {}
    for raw_line in text.splitlines():
        if not raw_line.lstrip().startswith("|"):
            continue
        cells = [clean_markdown_cell(cell) for cell in raw_line.strip().strip("|").split("|")]
        if len(cells) < 4 or not cells[0].isdigit():
            continue
        if cells[2].upper() != "RENAME":
            continue
        number = int(cells[0])
        old = cells[1]
        new = cells[3]
        if not old or not new:
            raise ValueError(f"incomplete rename row for chapter {number}")
        previous = result.get(number)
        pair = (old, new)
        if previous is not None and previous != pair:
            raise ValueError(f"conflicting audit rows for chapter {number}: {previous} vs {pair}")
        result[number] = pair
    return result


def parse_final_approval_overrides(text: str) -> dict[int, tuple[str, str]]:
    result: dict[int, tuple[str, str]] = {}
    for raw_line in text.splitlines():
        if not raw_line.lstrip().startswith("|"):
            continue
        cells = [clean_markdown_cell(cell) for cell in raw_line.strip().strip("|").split("|")]
        if len(cells) < 4 or not cells[0].isdigit():
            continue
        # FINAL_APPROVAL table: Ch | Current title | Final approved role title | confidence | ...
        number = int(cells[0])
        old = cells[1]
        new = cells[2]
        if old and new and normalize_title(old) != normalize_title(new):
            result[number] = (old, new)
    return result


def git_show(path: str, ref: str = AUDIT_REF) -> str:
    proc = subprocess.run(
        ["git", "show", f"{ref}:{path}"],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"cannot read {ref}:{path}: {proc.stderr.strip()}")
    return proc.stdout


def load_approved_renames() -> dict[int, tuple[str, str]]:
    renames: dict[int, tuple[str, str]] = {}
    for path in AUDIT_FILES:
        for number, pair in parse_audit_table(git_show(path)).items():
            previous = renames.get(number)
            if previous is not None and previous != pair:
                raise ValueError(f"conflicting frozen audit mapping for chapter {number}")
            renames[number] = pair

    for number, pair in parse_final_approval_overrides(git_show(FINAL_APPROVAL)).items():
        if number not in renames:
            raise ValueError(f"final approval override for non-rename chapter {number}")
        old, _ = renames[number]
        override_old, override_new = pair
        if normalize_title(old) != normalize_title(override_old):
            raise ValueError(
                f"final approval old-title mismatch for chapter {number}: {old!r} vs {override_old!r}"
            )
        renames[number] = (old, override_new)

    if RETIRED_CHAPTER in renames:
        raise ValueError("retired Chapter 150 must not enter the rename manifest")
    if len(renames) != EXPECTED_RENAMES:
        raise ValueError(f"expected {EXPECTED_RENAMES} approved renames, found {len(renames)}")
    if max(renames) > 488:
        raise ValueError("frozen audit must not silently cover post-488 chapters")
    return dict(sorted(renames.items()))


def _case_insensitive_replace(value: str, old: str, new: str) -> tuple[str, bool]:
    match = re.search(re.escape(old), value, flags=re.IGNORECASE)
    if not match:
        return value, False
    return value[: match.start()] + new + value[match.end() :], True


def extract_html_title(text: str) -> str | None:
    match = re.search(r"<h1[^>]*>(.*?)</h1>", text, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return None
    return html.unescape(re.sub(r"<[^>]+>", "", match.group(1))).strip()


def update_html_text(text: str, number: int, old: str, new: str) -> tuple[str, bool]:
    current = extract_html_title(text)
    if current is None:
        raise ValueError(f"chapter {number}: no h1 title found")
    if normalize_title(current) == normalize_title(new):
        return text, False
    if normalize_title(current) != normalize_title(old):
        raise ValueError(
            f"chapter {number}: current surfaced title {current!r} does not match audited old {old!r} "
            f"or approved new {new!r}"
        )

    original_article = text.split("<article", 1)[1] if "<article" in text else None

    def h1_repl(match: re.Match[str]) -> str:
        return f"{match.group(1)}{new.upper()}{match.group(3)}"

    updated, count = re.subn(
        r"(<h1[^>]*>)(.*?)(</h1>)",
        h1_repl,
        text,
        count=1,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if count != 1:
        raise ValueError(f"chapter {number}: failed to replace h1")

    def title_repl(match: re.Match[str]) -> str:
        inner, changed = _case_insensitive_replace(match.group(2), old, new)
        if not changed:
            inner, _ = _case_insensitive_replace(match.group(2), old.upper(), new)
        return f"{match.group(1)}{inner}{match.group(3)}"

    updated = re.sub(
        r"(<title>)(.*?)(</title>)",
        title_repl,
        updated,
        count=1,
        flags=re.IGNORECASE | re.DOTALL,
    )

    # Description metadata is a derivative surface; change only the audited title phrase.
    def meta_repl(match: re.Match[str]) -> str:
        whole = match.group(0)
        replacement, _ = _case_insensitive_replace(whole, old, new)
        return replacement

    updated = re.sub(
        r"<meta\b[^>]*\bname=[\"']description[\"'][^>]*>",
        meta_repl,
        updated,
        count=1,
        flags=re.IGNORECASE,
    )

    if original_article is not None:
        new_article = updated.split("<article", 1)[1]
        if new_article != original_article:
            raise ValueError(f"chapter {number}: HTML prose/article changed during title-only update")
    return updated, True


BOUNDARY_RE = re.compile(
    r"^(?:# CHAPTER (?P<standard>\d+)[ \t]*|## Chapter (?P<combined>\d+)[ \t]*[—–-][ \t]*(?P<combined_title>[^\n]+?)[ \t]*)$",
    flags=re.MULTILINE,
)


def markdown_titles(text: str) -> dict[int, str]:
    matches = list(BOUNDARY_RE.finditer(text))
    result: dict[int, str] = {}
    for i, match in enumerate(matches):
        number = int(match.group("standard") or match.group("combined"))
        combined = match.group("combined_title")
        if combined is not None:
            result[number] = combined.strip()
            continue
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk = text[match.end() : end]
        title_match = re.search(r"^##\s+(.+?)\s*$", chunk, flags=re.MULTILINE)
        if title_match:
            result[number] = title_match.group(1).strip()
    return result


def update_markdown_text(
    text: str,
    renames: dict[int, tuple[str, str]],
) -> tuple[str, set[int]]:
    matches = list(BOUNDARY_RE.finditer(text))
    replacements: list[tuple[int, int, str, int]] = []
    found: set[int] = set()

    for i, match in enumerate(matches):
        number = int(match.group("standard") or match.group("combined"))
        if number not in renames:
            continue
        old, new = renames[number]
        combined = match.group("combined_title")
        if combined is not None:
            current = combined.strip()
            found.add(number)
            if normalize_title(current) == normalize_title(new):
                continue
            if normalize_title(current) != normalize_title(old):
                raise ValueError(
                    f"chapter {number}: manuscript title {current!r} conflicts with audit {old!r} -> {new!r}"
                )
            replacements.append((match.start("combined_title"), match.end("combined_title"), new.upper(), number))
            continue

        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk_start = match.end()
        chunk = text[chunk_start:end]
        title_match = re.search(r"^##\s+(.+?)\s*$", chunk, flags=re.MULTILINE)
        if not title_match:
            raise ValueError(f"chapter {number}: no manuscript title heading after chapter boundary")
        current = title_match.group(1).strip()
        found.add(number)
        if normalize_title(current) == normalize_title(new):
            continue
        if normalize_title(current) != normalize_title(old):
            raise ValueError(
                f"chapter {number}: manuscript title {current!r} conflicts with audit {old!r} -> {new!r}"
            )
        start = chunk_start + title_match.start(1)
        stop = chunk_start + title_match.end(1)
        replacements.append((start, stop, new.upper(), number))

    updated = text
    changed: set[int] = set()
    for start, stop, replacement, number in reversed(replacements):
        updated = updated[:start] + replacement + updated[stop:]
        changed.add(number)
    return updated, changed


def update_markdown_file(path: Path, renames: dict[int, tuple[str, str]], expected: set[int]) -> set[int]:
    text = path.read_text(encoding="utf-8")
    titles = markdown_titles(text)
    missing = sorted(n for n in expected if n not in titles)
    if missing:
        raise ValueError(f"{path}: expected audited chapters missing: {missing}")
    updated, changed = update_markdown_text(text, {n: renames[n] for n in expected})
    if updated != text:
        path.write_text(updated, encoding="utf-8")
    return changed


def _replace_in_runs(paragraph, old: str, new: str) -> bool:
    runs = list(paragraph.runs)
    full = "".join(run.text for run in runs)
    match = re.search(re.escape(old), full, flags=re.IGNORECASE)
    if not match:
        return False
    start, stop = match.span()
    offset = 0
    first_written = False
    for run in runs:
        run_start = offset
        run_stop = offset + len(run.text)
        original = run.text
        offset = run_stop
        if run_stop <= start or run_start >= stop:
            continue
        left = original[: max(0, start - run_start)] if run_start <= start < run_stop else ""
        right = original[max(0, stop - run_start) :] if run_start < stop <= run_stop else ""
        if not first_written:
            run.text = left + new + right
            first_written = True
        else:
            run.text = right
    return True


def _docx_number(text: str) -> int | None:
    from scripts.promote_book1_polish import _number_from_heading

    return _number_from_heading(text)


def _docx_title_location(doc, chapter: int) -> tuple[int, object, str]:
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if _docx_number(paragraph.text) != chapter:
            continue
        lines = [line.strip() for line in paragraph.text.splitlines() if line.strip()]
        if len(lines) >= 2:
            return index, paragraph, lines[-1]
        for j in range(index + 1, len(paragraphs)):
            if _docx_number(paragraphs[j].text) is not None:
                break
            if paragraphs[j].text.strip():
                return j, paragraphs[j], paragraphs[j].text.strip()
        raise ValueError(f"DOCX chapter {chapter}: title paragraph not found")
    raise ValueError(f"DOCX chapter {chapter}: chapter heading not found")


def docx_titles(path: Path, chapters: Iterable[int]) -> dict[int, str]:
    from docx import Document

    doc = Document(path)
    return {chapter: _docx_title_location(doc, chapter)[2] for chapter in chapters}


def update_docx_titles(path: Path, renames: dict[int, tuple[str, str]], chapters: set[int]) -> set[int]:
    from docx import Document

    doc = Document(path)
    before = [paragraph.text for paragraph in doc.paragraphs]
    allowed_indices: set[int] = set()
    changed: set[int] = set()

    for chapter in sorted(chapters):
        old, new = renames[chapter]
        index, paragraph, current = _docx_title_location(doc, chapter)
        allowed_indices.add(index)
        if normalize_title(current) == normalize_title(new):
            continue
        if normalize_title(current) != normalize_title(old):
            raise ValueError(
                f"{path}: chapter {chapter} title {current!r} conflicts with audit {old!r} -> {new!r}"
            )
        if not _replace_in_runs(paragraph, current, new.upper()):
            raise ValueError(f"{path}: chapter {chapter} could not replace title in DOCX runs")
        changed.add(chapter)

    if not changed:
        return changed

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as tmp:
        tmp_path = Path(tmp.name)
    try:
        doc.save(tmp_path)
        check = Document(tmp_path)
        after = [paragraph.text for paragraph in check.paragraphs]
        if len(after) != len(before):
            raise ValueError(f"{path}: paragraph count changed during title integration")
        for index, (old_text, new_text) in enumerate(zip(before, after)):
            if index not in allowed_indices and old_text != new_text:
                raise ValueError(f"{path}: untargeted paragraph {index} changed")
        for chapter in chapters:
            current = _docx_title_location(check, chapter)[2]
            approved = renames[chapter][1]
            if normalize_title(current) != normalize_title(approved):
                raise ValueError(f"{path}: chapter {chapter} did not round-trip approved title")
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()
    return changed


def update_chapter_index(renames: dict[int, tuple[str, str]]) -> set[int]:
    text = CHAPTER_INDEX.read_text(encoding="utf-8")
    changed: set[int] = set()
    for number, (old, new) in renames.items():
        if number > 248:
            continue
        pattern = re.compile(rf"^{number}\. \*\*(.+?)\*\*$", flags=re.MULTILINE)
        match = pattern.search(text)
        if not match:
            raise ValueError(f"chapter index missing chapter {number}")
        current = match.group(1).strip()
        if normalize_title(current) == normalize_title(new):
            continue
        if normalize_title(current) != normalize_title(old):
            raise ValueError(
                f"chapter index {number}: {current!r} conflicts with audit {old!r} -> {new!r}"
            )
        replacement = f"{number}. **{new.upper()}**"
        text = text[: match.start()] + replacement + text[match.end() :]
        changed.add(number)
    if changed:
        CHAPTER_INDEX.write_text(text, encoding="utf-8")
    return changed


def source_route(number: int) -> Path:
    if number <= 82:
        return BOOK1_DOCX
    if number <= 137:
        return BOOK2_DOCX
    if number <= 155:
        return RUN_138_155
    if number <= 219:
        return RECOVERED_156_219
    if number <= 248:
        return RUNNING_220
    return Path(f"state/manuscript/Peg_Leg_Greg_Chapter_{number}_EXACT_WIP.md")


def reconcile_reader_surfaces(renames: dict[int, tuple[str, str]]) -> None:
    conflicts: list[str] = []
    for number, (old, new) in renames.items():
        path = Path("chapters") / f"{number:03d}.html"
        if not path.exists():
            conflicts.append(f"chapter {number}: missing {path}")
            continue
        current = extract_html_title(path.read_text(encoding="utf-8"))
        if current is None or normalize_title(current) not in {normalize_title(old), normalize_title(new)}:
            conflicts.append(f"chapter {number}: current title {current!r}, audit expects {old!r} or {new!r}")
    if conflicts:
        raise ValueError("current-authority reconciliation failed:\n" + "\n".join(conflicts))


def apply(renames: dict[int, tuple[str, str]]) -> None:
    reconcile_reader_surfaces(renames)
    changed_sources: set[int] = set()

    book1 = {n for n in renames if n <= 82}
    book2 = {n for n in renames if 83 <= n <= 137}
    md_138 = {n for n in renames if 138 <= n <= 155}
    md_156 = {n for n in renames if 156 <= n <= 219}
    md_220 = {n for n in renames if 220 <= n <= 248}
    checkpoints = {n for n in renames if n >= 249}

    if book1:
        changed_sources |= update_docx_titles(BOOK1_DOCX, renames, book1)
    if book2:
        changed_sources |= update_docx_titles(BOOK2_DOCX, renames, book2)
    if md_138:
        changed_sources |= update_markdown_file(RUN_138_155, renames, md_138)
    if md_156:
        changed_sources |= update_markdown_file(RECOVERED_156_219, renames, md_156)
    if md_220:
        changed_sources |= update_markdown_file(RUNNING_220, renames, md_220)
    for number in sorted(checkpoints):
        path = source_route(number)
        if not path.exists():
            raise ValueError(f"chapter {number}: exact checkpoint missing at {path}")
        changed_sources |= update_markdown_file(path, renames, {number})

    update_chapter_index(renames)

    changed_html: set[int] = set()
    for number, (old, new) in renames.items():
        for directory in (Path("chapters"), Path("light")):
            path = directory / f"{number:03d}.html"
            if not path.exists():
                continue
            text = path.read_text(encoding="utf-8")
            updated, changed = update_html_text(text, number, old, new)
            if changed:
                path.write_text(updated, encoding="utf-8")
                changed_html.add(number)

    print(
        f"approved renames={len(renames)}; canonical source titles changed={len(changed_sources)}; "
        f"reader/light title surfaces touched={len(changed_html)}"
    )


def _source_titles_for_verify(renames: dict[int, tuple[str, str]]) -> dict[int, str]:
    titles: dict[int, str] = {}
    book1 = sorted(n for n in renames if n <= 82)
    book2 = sorted(n for n in renames if 83 <= n <= 137)
    if book1:
        titles.update(docx_titles(BOOK1_DOCX, book1))
    if book2:
        titles.update(docx_titles(BOOK2_DOCX, book2))

    fixed_md = [RUN_138_155, RECOVERED_156_219, RUNNING_220]
    for path in fixed_md:
        if path.exists():
            titles.update({n: t for n, t in markdown_titles(path.read_text(encoding="utf-8")).items() if n in renames})
    for number in sorted(n for n in renames if n >= 249):
        path = source_route(number)
        if not path.exists():
            raise ValueError(f"verify: missing exact checkpoint {path}")
        parsed = markdown_titles(path.read_text(encoding="utf-8"))
        if number not in parsed:
            raise ValueError(f"verify: checkpoint {path} lacks chapter {number}")
        titles[number] = parsed[number]
    return titles


def verify(renames: dict[int, tuple[str, str]]) -> None:
    source_titles = _source_titles_for_verify(renames)
    problems: list[str] = []
    for number, (_, new) in renames.items():
        source = source_titles.get(number)
        if source is None:
            problems.append(f"chapter {number}: canonical source title unavailable")
        elif normalize_title(source) != normalize_title(new):
            problems.append(f"chapter {number}: canonical source {source!r} != approved {new!r}")

        chapter_path = Path("chapters") / f"{number:03d}.html"
        if not chapter_path.exists():
            problems.append(f"chapter {number}: illustrated surface missing")
        else:
            surfaced = extract_html_title(chapter_path.read_text(encoding="utf-8"))
            if surfaced is None or normalize_title(surfaced) != normalize_title(new):
                problems.append(f"chapter {number}: illustrated title {surfaced!r} != approved {new!r}")

        light_path = Path("light") / f"{number:03d}.html"
        if light_path.exists():
            surfaced = extract_html_title(light_path.read_text(encoding="utf-8"))
            if surfaced is None or normalize_title(surfaced) != normalize_title(new):
                problems.append(f"chapter {number}: Text Reader title {surfaced!r} != approved {new!r}")

    if CHAPTER_INDEX.exists():
        index_text = CHAPTER_INDEX.read_text(encoding="utf-8")
        for number, (_, new) in renames.items():
            if number > 248:
                continue
            match = re.search(rf"^{number}\. \*\*(.+?)\*\*$", index_text, flags=re.MULTILINE)
            current = match.group(1) if match else None
            if current is None or normalize_title(current) != normalize_title(new):
                problems.append(f"chapter index {number}: {current!r} != approved {new!r}")

    manifest = Path("light/manifest.json")
    if manifest.exists():
        data = json.loads(manifest.read_text(encoding="utf-8"))
        manifest_titles = {int(item["number"]): str(item["title"]) for item in data.get("chapters", [])}
        for number, (_, new) in renames.items():
            if number in manifest_titles and normalize_title(manifest_titles[number]) != normalize_title(new):
                problems.append(f"light manifest {number}: {manifest_titles[number]!r} != approved {new!r}")

    exports_dir = Path("exports/notebooklm")
    if exports_dir.exists():
        export_titles: dict[int, str] = {}
        for path in exports_dir.glob("PLG_BOOK_*.md"):
            export_titles.update(markdown_titles(path.read_text(encoding="utf-8")))
        for number, (_, new) in renames.items():
            if number in export_titles and normalize_title(export_titles[number]) != normalize_title(new):
                problems.append(f"NotebookLM export {number}: {export_titles[number]!r} != approved {new!r}")

    anchors = {
        5: "The Warrior",
        177: "The Stagehand",
        231: "The Magistrate",
        331: "The Surveyor",
        446: "The Investor",
    }
    for number, expected in anchors.items():
        path = Path("chapters") / f"{number:03d}.html"
        current = extract_html_title(path.read_text(encoding="utf-8")) if path.exists() else None
        if current is None or normalize_title(current) != normalize_title(expected):
            problems.append(f"anchor chapter {number}: {current!r} != {expected!r}")

    if problems:
        raise ValueError("title integration verification failed:\n" + "\n".join(problems))
    print(f"verified {len(renames)} approved role-title renames across canonical and reader surfaces")


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--apply", action="store_true")
    mode.add_argument("--verify", action="store_true")
    args = parser.parse_args()

    renames = load_approved_renames()
    if args.apply:
        apply(renames)
    else:
        verify(renames)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
